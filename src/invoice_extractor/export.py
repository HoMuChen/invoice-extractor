"""Write extraction results to an Excel workbook (Invoices / Line Items / Review) and JSON."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .extractor import Extraction

INVOICE_COLUMNS = [
    "file", "status", "invoice_number", "invoice_date", "due_date", "vendor_name", "vendor_address",
    "customer_name", "currency", "line_count", "subtotal", "tax", "shipping", "total", "issues",
]
LINE_COLUMNS = ["file", "invoice_number", "line", "description", "quantity", "unit_price", "amount"]
STATUS_FILL = {
    "OK": PatternFill("solid", fgColor="E2F0D9"),
    "REVIEW": PatternFill("solid", fgColor="FFF2CC"),
    "ERROR": PatternFill("solid", fgColor="F8CBAD"),
}


def invoice_row(ex: Extraction) -> dict:
    inv = ex.invoice
    row = {"file": ex.path.name, "status": ex.status, "issues": "; ".join(ex.issues) or ex.error}
    if inv:
        row.update(inv.model_dump(exclude={"line_items"}))
        row["line_count"] = len(inv.line_items)
    return row


def write_xlsx(results: list[Extraction], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    ws = wb.active
    ws.title = "Invoices"
    ws.append(INVOICE_COLUMNS)
    for ex in results:
        row = invoice_row(ex)
        ws.append([row.get(c) for c in INVOICE_COLUMNS])
        ws.cell(row=ws.max_row, column=2).fill = STATUS_FILL[ex.status]

    lines = wb.create_sheet("Line Items")
    lines.append(LINE_COLUMNS)
    for ex in results:
        if not ex.invoice:
            continue
        for i, item in enumerate(ex.invoice.line_items, start=1):
            lines.append([ex.path.name, ex.invoice.invoice_number, i, item.description,
                          item.quantity, item.unit_price, item.amount])

    review = wb.create_sheet("Review")
    review.append(["file", "status", "issue"])
    for ex in results:
        for issue in ex.issues or ([ex.error] if ex.error else []):
            review.append([ex.path.name, ex.status, issue])

    for sheet in (ws, lines, review):
        _style(sheet)
    wb.save(path)


def write_json(results: list[Extraction], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = [
        {"file": ex.path.name, "status": ex.status, "issues": ex.issues, "error": ex.error,
         "invoice": ex.invoice.model_dump() if ex.invoice else None}
        for ex in results
    ]
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=False))


def _style(ws) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for idx in range(1, ws.max_column + 1):
        letter = get_column_letter(idx)
        width = max(len(str(c.value or "")) for c in ws[letter][:200])
        ws.column_dimensions[letter].width = min(width + 2, 60)
