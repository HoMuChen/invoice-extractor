import json
from pathlib import Path

from openpyxl import load_workbook

from invoice_extractor.evaluate import evaluate
from invoice_extractor.export import write_json, write_xlsx
from invoice_extractor.extractor import Extraction


def results(invoice) -> list[Extraction]:
    bad = invoice.model_copy(deep=True)
    bad.total = 2000.0
    return [
        Extraction(path=Path("ok.pdf"), invoice=invoice),
        Extraction(path=Path("review.pdf"), invoice=bad, issues=["Subtotal + tax + shipping = 1556.09, total printed 2000.00"]),
        Extraction(path=Path("broken.pdf"), error="Network error"),
    ]


def test_write_xlsx(tmp_path, invoice):
    write_xlsx(results(invoice), tmp_path / "out.xlsx")
    wb = load_workbook(tmp_path / "out.xlsx")
    assert wb.sheetnames == ["Invoices", "Line Items", "Review"]
    inv = wb["Invoices"]
    assert [inv.cell(row=r, column=2).value for r in range(2, 5)] == ["OK", "REVIEW", "ERROR"]
    assert wb["Line Items"].max_row == 1 + 2 * 2  # two invoices with two lines each
    review_rows = [row for row in wb["Review"].iter_rows(min_row=2, values_only=True)]
    assert review_rows == [("review.pdf", "REVIEW", "Subtotal + tax + shipping = 1556.09, total printed 2000.00"),
                           ("broken.pdf", "ERROR", "Network error")]


def test_write_json(tmp_path, invoice):
    write_json(results(invoice), tmp_path / "out.json")
    data = json.loads((tmp_path / "out.json").read_text())
    assert [d["status"] for d in data] == ["OK", "REVIEW", "ERROR"]
    assert data[0]["invoice"]["invoice_number"] == "INV-1001" and data[2]["invoice"] is None


def test_evaluate_against_truth(tmp_path, invoice):
    (tmp_path / "ok.json").write_text(invoice.model_dump_json())
    (tmp_path / "review.json").write_text(invoice.model_dump_json())
    score = evaluate(results(invoice)[:2], tmp_path)
    fields_per_invoice = 10 + 3 * 2
    assert score.total == 2 * fields_per_invoice
    assert score.correct == score.total - 1
    assert score.mismatches == ["review.pdf total: got 2000.0, expected 1556.09"]
